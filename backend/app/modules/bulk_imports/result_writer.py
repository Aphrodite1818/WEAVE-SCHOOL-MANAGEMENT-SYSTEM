# ================================ #
#   bulk_imports_result_writer.py  #
# ================================ #

"""Result file writers for tenant bulk import workflows."""

from __future__ import annotations

import csv
import html
import io
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.modules.bulk_imports.models import ImportResourceType

SPREADSHEET_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HTML_CONTENT_TYPE = "text/html; charset=utf-8"


@dataclass(frozen=True)
class ImportResultFile:
    """Generated import result file."""

    filename: str
    content_type: str
    content_bytes: bytes


def create_result_filename(
    *,
    resource_type: ImportResourceType,
    suffix: str,
    extension: str = "xlsx",
) -> str:
    """Create a timestamped result filename."""

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return f"{resource_type.value}_{suffix}_{timestamp}.{extension}"


def collect_csv_headers(
    *,
    rows: list[dict[str, Any]],
    preferred_headers: list[str] | None = None,
) -> list[str]:
    """Collect stable report headers from row dictionaries."""

    headers: list[str] = []

    for header in preferred_headers or []:
        if header not in headers:
            headers.append(header)

    for row in rows:
        for key in row.keys():
            if key not in headers:
                headers.append(key)

    return headers


def convert_value_for_csv(value: Any) -> str:
    """Convert a Python value into a spreadsheet-safe string."""

    if value is None:
        return ""

    if isinstance(value, dict | list):
        return str(value)

    return str(value)


def write_csv_bytes(
    *,
    rows: list[dict[str, Any]],
    headers: list[str],
) -> bytes:
    """Write rows to CSV bytes."""

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
    writer.writeheader()

    for row in rows:
        writer.writerow(
            {
                header: convert_value_for_csv(row.get(header))
                for header in headers
            }
        )

    return output.getvalue().encode("utf-8-sig")


def write_xlsx_bytes(
    *,
    rows: list[dict[str, Any]],
    headers: list[str],
    title: str = "Import Result",
) -> bytes:
    """Write result rows to a styled XLSX workbook."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(bottom=Side(style="thin", color="E5E7EB"))

    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row in rows:
        sheet.append([convert_value_for_csv(row.get(header)) for header in headers])

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_index, header in enumerate(headers, start=1):
        values = [convert_value_for_csv(row.get(header)) for row in rows[:250]]
        max_length = max([len(str(header)), *[len(value) for value in values]], default=len(str(header)))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 42)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def create_result_report(
    *,
    resource_type: ImportResourceType,
    result_rows: list[dict[str, Any]],
) -> ImportResultFile:
    """Create an XLSX report containing successful and failed row outcomes."""

    preferred_headers_by_resource = {
        ImportResourceType.STUDENTS: [
            "row_number",
            "status",
            "first_name",
            "last_name",
            "admission_number",
            "class_name",
            "class_arm",
            "setup_code",
            "access_code_expires_at",
            "parent_invitations_queued",
            "error_message",
        ],
    }

    headers = collect_csv_headers(
        rows=result_rows,
        preferred_headers=preferred_headers_by_resource.get(resource_type, []),
    )

    if not headers:
        headers = ["row_number", "status", "error_message"]

    return ImportResultFile(
        filename=create_result_filename(resource_type=resource_type, suffix="result", extension="xlsx"),
        content_type=SPREADSHEET_CONTENT_TYPE,
        content_bytes=write_xlsx_bytes(
            rows=result_rows,
            headers=headers,
            title=f"{resource_type.value}_result",
        ),
    )


def _format_slip_value(value: Any) -> str:
    """Format a value for safe HTML display."""

    return html.escape(convert_value_for_csv(value) or "--")


def _student_full_name(row: dict[str, Any]) -> str:
    """Build a student display name from a result row."""

    name = " ".join(
        part
        for part in [
            convert_value_for_csv(row.get("first_name")).strip(),
            convert_value_for_csv(row.get("last_name")).strip(),
        ]
        if part
    )
    return name or "Student"


def _student_class_display(row: dict[str, Any]) -> str:
    """Build a class display string from a result row."""

    class_name = convert_value_for_csv(row.get("class_name")).strip()
    class_arm = convert_value_for_csv(row.get("class_arm")).strip()
    return " ".join(part for part in [class_name, class_arm] if part) or "--"


def _student_slip_rows(result_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Return student rows that have generated setup codes."""

    return [
        row
        for row in result_rows
        if str(row.get("status") or "").lower() == "created" and convert_value_for_csv(row.get("setup_code")).strip()
    ]


def create_student_access_slip_report(
    *,
    result_rows: list[dict[str, Any]],
    school_name: str | None,
) -> ImportResultFile:
    """Create a printable HTML slip report for imported student setup codes."""

    slip_rows = _student_slip_rows(result_rows)
    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    safe_school_name = html.escape(school_name or "School")

    if slip_rows:
        slip_markup = "\n".join(
            f"""
            <article class="slip">
              <header class="slip-header">
                <div>
                  <p class="eyebrow">Student Access Slip</p>
                  <h2>{html.escape(_student_full_name(row))}</h2>
                </div>
                <span class="badge">Initial Setup</span>
              </header>
              <table>
                <tbody>
                  <tr><th>Student Name</th><td>{html.escape(_student_full_name(row))}</td></tr>
                  <tr><th>Admission Number</th><td class="strong">{_format_slip_value(row.get("admission_number"))}</td></tr>
                  <tr><th>Class</th><td>{html.escape(_student_class_display(row))}</td></tr>
                  <tr><th>Setup / Access Code</th><td class="code">{_format_slip_value(row.get("setup_code"))}</td></tr>
                  <tr><th>Code Expires</th><td>{_format_slip_value(row.get("access_code_expires_at"))}</td></tr>
                </tbody>
              </table>
              <p class="note">Give this slip only to the student or guardian. The student must change their password after first login.</p>
            </article>
            """
            for row in slip_rows
        )
    else:
        slip_markup = """
        <section class="empty-state">
          <h2>No printable student slips available</h2>
          <p>This report only includes successfully created student rows that have generated setup codes.</p>
        </section>
        """

    document = f"""
    <!doctype html>
    <html lang="en">
      <head>
        <meta charset="utf-8" />
        <meta name="viewport" content="width=device-width, initial-scale=1" />
        <title>{safe_school_name} Student Access Slips</title>
        <style>
          :root {{
            color-scheme: light;
            --ink: #111827;
            --muted: #6b7280;
            --line: #d1d5db;
            --soft: #f3f4f6;
            --brand: #2563eb;
            --brand-soft: #eff6ff;
          }}
          * {{ box-sizing: border-box; }}
          body {{
            margin: 0;
            background: #f8fafc;
            color: var(--ink);
            font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, "Segoe UI", sans-serif;
          }}
          .toolbar {{
            position: sticky;
            top: 0;
            z-index: 10;
            display: flex;
            justify-content: space-between;
            gap: 1rem;
            padding: 1rem 1.25rem;
            border-bottom: 1px solid var(--line);
            background: rgba(255, 255, 255, 0.94);
            backdrop-filter: blur(12px);
          }}
          .toolbar h1 {{ margin: 0; font-size: 1rem; }}
          .toolbar p {{ margin: 0.25rem 0 0; color: var(--muted); font-size: 0.8rem; }}
          button {{
            border: 0;
            border-radius: 999px;
            background: var(--brand);
            color: white;
            cursor: pointer;
            font-weight: 700;
            padding: 0.75rem 1rem;
          }}
          main {{ padding: 1rem; }}
          .sheet {{
            display: grid;
            grid-template-columns: repeat(2, minmax(0, 1fr));
            gap: 0.75rem;
            max-width: 1100px;
            margin: 0 auto;
          }}
          .slip {{
            break-inside: avoid;
            overflow: hidden;
            border: 1px solid var(--line);
            border-radius: 18px;
            background: white;
            box-shadow: 0 12px 30px rgba(15, 23, 42, 0.08);
          }}
          .slip-header {{
            display: flex;
            align-items: flex-start;
            justify-content: space-between;
            gap: 1rem;
            padding: 1rem;
            border-bottom: 1px solid var(--line);
            background: linear-gradient(135deg, var(--brand-soft), #ffffff);
          }}
          .eyebrow {{
            margin: 0 0 0.25rem;
            color: var(--brand);
            font-size: 0.72rem;
            font-weight: 800;
            letter-spacing: 0.08em;
            text-transform: uppercase;
          }}
          h2 {{ margin: 0; font-size: 1.05rem; }}
          .badge {{
            white-space: nowrap;
            border-radius: 999px;
            background: var(--brand);
            color: white;
            font-size: 0.72rem;
            font-weight: 800;
            padding: 0.35rem 0.65rem;
          }}
          table {{ width: 100%; border-collapse: collapse; }}
          th, td {{
            border-bottom: 1px solid var(--line);
            padding: 0.68rem 0.85rem;
            text-align: left;
            vertical-align: top;
            font-size: 0.86rem;
          }}
          th {{ width: 42%; background: var(--soft); color: #374151; font-weight: 800; }}
          td {{ font-weight: 600; }}
          .strong {{ font-weight: 900; letter-spacing: 0.02em; }}
          .code {{
            color: var(--brand);
            font-family: ui-monospace, SFMono-Regular, Menlo, Monaco, Consolas, "Liberation Mono", monospace;
            font-size: 1.2rem;
            font-weight: 900;
            letter-spacing: 0.12em;
          }}
          .note {{
            margin: 0;
            padding: 0.8rem 1rem 1rem;
            color: var(--muted);
            font-size: 0.76rem;
            line-height: 1.45;
          }}
          .empty-state {{
            grid-column: 1 / -1;
            border: 1px dashed var(--line);
            border-radius: 20px;
            background: white;
            padding: 2rem;
            text-align: center;
          }}
          @media print {{
            @page {{ size: A4; margin: 10mm; }}
            body {{ background: white; }}
            .toolbar {{ display: none; }}
            main {{ padding: 0; }}
            .sheet {{ gap: 6mm; max-width: none; }}
            .slip {{ box-shadow: none; border-radius: 12px; }}
          }}
          @media (max-width: 760px) {{
            .sheet {{ grid-template-columns: 1fr; }}
            .toolbar {{ align-items: flex-start; flex-direction: column; }}
          }}
        </style>
      </head>
      <body>
        <div class="toolbar">
          <div>
            <h1>{safe_school_name} — Student Access Slips</h1>
            <p>{len(slip_rows)} printable slips · Generated {generated_at}</p>
          </div>
          <button type="button" onclick="window.print()">Print slips</button>
        </div>
        <main>
          <section class="sheet">
            {slip_markup}
          </section>
        </main>
      </body>
    </html>
    """

    return ImportResultFile(
        filename=create_result_filename(
            resource_type=ImportResourceType.STUDENTS,
            suffix="access_slips",
            extension="html",
        ),
        content_type=HTML_CONTENT_TYPE,
        content_bytes=document.encode("utf-8"),
    )


def create_error_report(
    *,
    resource_type: ImportResourceType,
    row_errors: list[dict[str, Any]],
) -> ImportResultFile:
    """Create a CSV report containing only row errors."""

    headers = collect_csv_headers(
        rows=row_errors,
        preferred_headers=[
            "row_number",
            "field_name",
            "error_code",
            "error_message",
        ],
    )

    return ImportResultFile(
        filename=create_result_filename(resource_type=resource_type, suffix="errors", extension="csv"),
        content_type="text/csv",
        content_bytes=write_csv_bytes(rows=row_errors, headers=headers),
    )
