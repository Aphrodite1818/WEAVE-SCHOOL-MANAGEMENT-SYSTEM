# ================================== #
#   bulk_imports_template_writer.py  #
# ================================== #

"""Generate downloadable XLSX bulk import templates."""

from __future__ import annotations

import io
from dataclasses import dataclass
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.bulk_imports.models import ImportFileType, ImportResourceType
from app.modules.bulk_imports.template_security import (
    build_headers_hash,
    build_template_signature_payload,
    sign_template_payload,
)
from app.modules.bulk_imports.templates import (
    DATA_HEADERS_BY_RESOURCE,
    TEMPLATE_METADATA_SHEET_NAME,
    TEMPLATE_VERSION_BY_RESOURCE,
    get_template_definition,
)


@dataclass(frozen=True)
class GeneratedImportTemplate:
    """Generated import template file payload."""

    filename: str
    content_bytes: bytes
    content_type: str


def build_template_metadata(
    *,
    tenant_id: UUID,
    resource_type: ImportResourceType,
) -> dict[str, str]:
    """Build signed metadata for a tenant-specific import template."""

    data_headers = DATA_HEADERS_BY_RESOURCE[resource_type]
    template_version = TEMPLATE_VERSION_BY_RESOURCE[resource_type]
    headers_hash = build_headers_hash(data_headers)
    signature_payload = build_template_signature_payload(
        tenant_id=tenant_id,
        resource_type=resource_type,
        template_version=template_version,
        headers_hash=headers_hash,
    )

    return {
        "_import_resource_type": resource_type.value,
        "_import_template_version": template_version,
        "_import_headers_hash": headers_hash,
        "_import_template_signature": sign_template_payload(signature_payload),
    }


def create_xlsx_template(
    *,
    tenant_id: UUID,
    resource_type: ImportResourceType,
) -> GeneratedImportTemplate:
    """Create an XLSX template with hidden signed metadata."""

    workbook = Workbook()

    data_sheet = workbook.active
    data_sheet.title = f"{resource_type.value}_import"

    data_headers = DATA_HEADERS_BY_RESOURCE[resource_type]
    data_sheet.append(data_headers)
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = data_sheet.dimensions

    header_fill = PatternFill("solid", fgColor="EAF2FF")
    for column_index, header in enumerate(data_headers, start=1):
        cell = data_sheet.cell(row=1, column=column_index)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        data_sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(len(header) + 6, 16), 34
        )

    template_definition = get_template_definition(resource_type=resource_type)
    instructions_sheet = workbook.create_sheet("Instructions")
    instructions_sheet.append(["Student import instructions"])
    instructions_sheet["A1"].font = Font(bold=True, size=14)
    for note in template_definition.notes:
        instructions_sheet.append([note])
    instructions_sheet.append([""])
    instructions_sheet.append(
        ["Column", "Required", "Example", "Accepted values", "Description"]
    )
    header_row_number = instructions_sheet.max_row
    for cell in instructions_sheet[header_row_number]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column in template_definition.columns:
        instructions_sheet.append(
            [
                column.name,
                "Yes" if column.required else "No",
                column.example or "",
                ", ".join(column.accepted_values),
                column.description or "",
            ]
        )
    instructions_sheet.column_dimensions["A"].width = 30
    instructions_sheet.column_dimensions["B"].width = 12
    instructions_sheet.column_dimensions["C"].width = 28
    instructions_sheet.column_dimensions["D"].width = 36
    instructions_sheet.column_dimensions["E"].width = 80

    metadata = build_template_metadata(
        tenant_id=tenant_id,
        resource_type=resource_type,
    )

    metadata_sheet = workbook.create_sheet(TEMPLATE_METADATA_SHEET_NAME)
    metadata_sheet.sheet_state = "veryHidden"
    metadata_sheet.append(["key", "value"])

    for key, value in metadata.items():
        metadata_sheet.append([key, value])

    output = io.BytesIO()
    workbook.save(output)

    return GeneratedImportTemplate(
        filename=f"{resource_type.value}_import_template.xlsx",
        content_bytes=output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def create_import_template(
    *,
    tenant_id: UUID,
    resource_type: ImportResourceType,
    file_type: ImportFileType,
) -> GeneratedImportTemplate:
    """Create a signed XLSX import template file."""

    if file_type != ImportFileType.XLSX:
        raise ValueError(
            "Bulk import templates are XLSX-only. Download the .xlsx template and upload that file."
        )

    return create_xlsx_template(tenant_id=tenant_id, resource_type=resource_type)
