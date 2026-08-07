# ========================== #
#   bulk_imports_parsers.py  #
# ========================== #

"""XLSX parser for tenant bulk import uploads."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path, PurePosixPath
from typing import Any, Sequence
from zipfile import BadZipFile, ZipFile

from fastapi import UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from app.core.exceptions import ImportParserError
from app.modules.bulk_imports.models import ImportFileType
from app.modules.bulk_imports.templates import TEMPLATE_METADATA_SHEET_NAME

MAX_IMPORT_FILE_SIZE_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5_000
MAX_XLSX_ARCHIVE_ENTRIES = 2_000
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 200
MAX_XLSX_WORKSHEET_ROWS = 10_000
MAX_XLSX_WORKSHEET_COLUMNS = 64


@dataclass(frozen=True)
class ParsedImportRow:
    """One parsed row from an uploaded import file."""

    row_number: int
    raw_data: dict[str, Any]


@dataclass(frozen=True)
class ParsedImportFile:
    """Parsed import file plus file-level metadata."""

    file_type: ImportFileType
    file_size_bytes: int
    headers: list[str]
    rows: list[ParsedImportRow]
    metadata: dict[str, Any] = field(default_factory=dict)


def _clean_headers(headers: Sequence[Any]) -> list[str]:
    """Normalize parser-level headers by trimming whitespace."""

    cleaned_headers: list[str] = []
    for header in headers:
        if header is None:
            cleaned_headers.append("")
            continue
        cleaned_headers.append(str(header).strip())
    return cleaned_headers


def _trim_trailing_blank_headers(headers: list[str]) -> list[str]:
    """Remove Excel artifact columns that appear after the real header set."""

    trimmed_headers = list(headers)
    while trimmed_headers and not trimmed_headers[-1]:
        trimmed_headers.pop()
    return trimmed_headers


def _has_values_beyond_headers(*, headers: list[str], values: Sequence[Any]) -> bool:
    """Return True if row data exists beyond the accepted header range."""

    if len(values) <= len(headers):
        return False
    for value in values[len(headers) :]:
        serialized_value = _to_serializable_value(value)
        if serialized_value is not None and str(serialized_value).strip() != "":
            return True
    return False


def _ensure_unique_headers(headers: list[str]) -> None:
    """Reject duplicate non-empty headers."""

    seen_headers: set[str] = set()
    for header in headers:
        if not header:
            continue
        normalized_header = header.strip().lower()
        if normalized_header in seen_headers:
            raise ImportParserError(f"Duplicate column header found: {header}")
        seen_headers.add(normalized_header)


def _validate_xlsx_archive(file_bytes: bytes) -> None:
    """Reject suspicious ZIP containers before openpyxl expands them."""

    try:
        with ZipFile(io.BytesIO(file_bytes)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
                raise ImportParserError(
                    "XLSX archive contains too many internal files."
                )

            total_uncompressed = 0
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if path.is_absolute() or ".." in path.parts:
                    raise ImportParserError(
                        "XLSX archive contains an unsafe internal path."
                    )
                if entry.flag_bits & 0x1:
                    raise ImportParserError(
                        "Encrypted XLSX archives are not supported."
                    )

                total_uncompressed += int(entry.file_size or 0)
                if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise ImportParserError(
                        "XLSX archive expands beyond the allowed size."
                    )

                compressed_size = max(int(entry.compress_size or 0), 1)
                if entry.file_size / compressed_size > MAX_XLSX_COMPRESSION_RATIO:
                    raise ImportParserError(
                        "XLSX archive has a suspicious compression ratio."
                    )
    except BadZipFile as exc:
        raise ImportParserError(
            "Uploaded file is not a readable XLSX workbook. Download a fresh backend-generated template and try again."
        ) from exc


def _read_xlsx_metadata(workbook) -> dict[str, Any]:
    """Read template metadata from the hidden XLSX metadata sheet."""

    if TEMPLATE_METADATA_SHEET_NAME not in workbook.sheetnames:
        return {}
    worksheet = workbook[TEMPLATE_METADATA_SHEET_NAME]
    metadata: dict[str, Any] = {}
    for row in worksheet.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        key = _to_serializable_value(row[0])
        value = _to_serializable_value(row[1])
        if not key or str(key).strip().lower() == "key":
            continue
        metadata[str(key).strip()] = value
    return metadata


def _select_xlsx_data_sheet(workbook) -> Worksheet:
    """Select the first visible non-metadata worksheet as the import data sheet."""

    for worksheet in workbook.worksheets:
        if worksheet.title == TEMPLATE_METADATA_SHEET_NAME:
            continue
        if getattr(worksheet, "sheet_state", "visible") == "visible":
            return worksheet
    raise ImportParserError("XLSX file must contain a visible import data sheet.")


def _validate_worksheet_dimensions(worksheet) -> None:
    """Reject worksheets whose declared dimensions could exhaust resources."""

    if worksheet.max_row and worksheet.max_row > MAX_XLSX_WORKSHEET_ROWS:
        raise ImportParserError(
            f"XLSX worksheet cannot exceed {MAX_XLSX_WORKSHEET_ROWS} rows including blank/formatted rows."
        )
    if worksheet.max_column and worksheet.max_column > MAX_XLSX_WORKSHEET_COLUMNS:
        raise ImportParserError(
            f"XLSX worksheet cannot exceed {MAX_XLSX_WORKSHEET_COLUMNS} columns."
        )


def _is_blank_row(row_data: dict[str, Any]) -> bool:
    """Return True when all row values are blank."""

    return all(value is None or str(value).strip() == "" for value in row_data.values())


def _to_serializable_value(value: Any) -> Any:
    """Convert spreadsheet values into JSON-safe values."""

    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        cleaned_value = value.strip()
        return cleaned_value or None
    return value


def _build_row_from_values(
    *,
    headers: list[str],
    values: Sequence[Any],
) -> dict[str, Any]:
    """Build one raw row dictionary from headers and row values."""

    return {
        headers[index]: _to_serializable_value(value)
        for index, value in enumerate(values)
        if index < len(headers) and headers[index]
    }


class BulkImportParser:
    """Parse backend-generated XLSX files into row dictionaries."""

    @staticmethod
    def resolve_file_type(
        *,
        filename: str | None,
        file_type: ImportFileType | None = None,
    ) -> ImportFileType:
        """Resolve and enforce XLSX file type."""

        if file_type is not None and file_type != ImportFileType.XLSX:
            raise ImportParserError(
                "Bulk imports are XLSX-only. Download the backend-generated .xlsx template."
            )
        if not filename:
            raise ImportParserError(
                "Uploaded file must have a filename ending in .xlsx."
            )
        extension = Path(filename).suffix.lower().lstrip(".")
        if extension != ImportFileType.XLSX.value:
            raise ImportParserError(
                "Bulk imports only support .xlsx files. Download the backend-generated XLSX template and upload it without converting it."
            )
        return ImportFileType.XLSX

    @staticmethod
    def parse_xlsx_bytes(
        *,
        file_bytes: bytes,
        file_size_bytes: int,
        max_rows: int = MAX_IMPORT_ROWS,
    ) -> ParsedImportFile:
        """Parse XLSX bytes into a parsed import file."""

        _validate_xlsx_archive(file_bytes)
        workbook = None
        try:
            workbook = load_workbook(
                filename=io.BytesIO(file_bytes),
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            worksheet = _select_xlsx_data_sheet(workbook)
            _validate_worksheet_dimensions(worksheet)
            metadata = _read_xlsx_metadata(workbook)
            rows_iter = worksheet.iter_rows(values_only=True)

            try:
                header_row = next(rows_iter)
            except StopIteration as exc:
                raise ImportParserError("XLSX file is empty.") from exc

            headers = _trim_trailing_blank_headers(_clean_headers(header_row))
            if not any(headers):
                raise ImportParserError(
                    "XLSX file must contain at least one valid column header."
                )
            if len(headers) > MAX_XLSX_WORKSHEET_COLUMNS:
                raise ImportParserError(
                    f"XLSX worksheet cannot exceed {MAX_XLSX_WORKSHEET_COLUMNS} columns."
                )
            _ensure_unique_headers(headers)

            parsed_rows: list[ParsedImportRow] = []
            for row_number, row_values in enumerate(rows_iter, start=2):
                if row_number > MAX_XLSX_WORKSHEET_ROWS:
                    raise ImportParserError(
                        "XLSX worksheet exceeds the safe row scan limit."
                    )
                if _has_values_beyond_headers(headers=headers, values=row_values):
                    raise ImportParserError(
                        f"Row {row_number} contains data outside the template columns. Remove extra columns and try again."
                    )

                raw_data = _build_row_from_values(headers=headers, values=row_values)
                if _is_blank_row(raw_data):
                    continue

                parsed_rows.append(
                    ParsedImportRow(row_number=row_number, raw_data=raw_data)
                )
                if len(parsed_rows) > max_rows:
                    raise ImportParserError(
                        f"Import file cannot exceed {max_rows} data rows."
                    )

            return ParsedImportFile(
                file_type=ImportFileType.XLSX,
                file_size_bytes=file_size_bytes,
                headers=headers,
                rows=parsed_rows,
                metadata=metadata,
            )
        except ImportParserError:
            raise
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise ImportParserError(
                "Uploaded file is not a readable XLSX workbook. Download a fresh backend-generated template and try again."
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    async def parse_upload(
        upload_file: UploadFile,
        *,
        file_type: ImportFileType | None = None,
        max_file_size_bytes: int = MAX_IMPORT_FILE_SIZE_BYTES,
        max_rows: int = MAX_IMPORT_ROWS,
    ) -> ParsedImportFile:
        """Read an uploaded XLSX file and return parsed file metadata and rows."""

        file_bytes = await upload_file.read(max_file_size_bytes + 1)
        if not file_bytes:
            raise ImportParserError("Uploaded file is empty.")

        file_size_bytes = len(file_bytes)
        if file_size_bytes > max_file_size_bytes:
            raise ImportParserError(
                f"Uploaded file is too large. Maximum allowed size is {max_file_size_bytes} bytes."
            )

        BulkImportParser.resolve_file_type(
            filename=upload_file.filename, file_type=file_type
        )
        return BulkImportParser.parse_xlsx_bytes(
            file_bytes=file_bytes,
            file_size_bytes=file_size_bytes,
            max_rows=max_rows,
        )
