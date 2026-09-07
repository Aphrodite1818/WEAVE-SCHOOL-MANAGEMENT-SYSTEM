"""
Generates fake student data as an Excel (.xlsx) file, using the Faker library.

Academic placement is expressed as two independent fields — level and
arm — matching Weave's current student import contract. Full class
display names like "JSS1 A" or "SS1 B" are NOT written anywhere; the
import template expects level/arm as separate columns instead.

Columns (in this order):
first_name, last_name, date_of_birth, gender, level, arm,
state_of_origin, parent_email_1, parent_relationship_1, parent_email_2,
parent_relationship_2

Class / arm rules (from the school's actual data):
- Levels: JSS1, JSS2, JSS3, SS1, SS2, SS3
- Arms: A, B, C for every level (18 classes total)

Department is no longer part of the import contract, so this script
does not generate or write a department column at all.

Parent email coverage:
- We do NOT want every student to have a parent email — just enough to
  guarantee at least one student per class (level + arm combination) has
  one, so every class shows up with at least one parent contact.
- All parent emails are aliases of a single Gmail inbox using the
  "+" trick (e.g. dfragraid+parent1@gmail.com, dfragraid+parent2@gmail.com,
  ...). Gmail ignores everything after "+" up to "@", so these all land
  in the same inbox but read as distinct addresses to any system that
  stores them.
- Each covered student gets exactly ONE parent email (parent_email_1 +
  parent_relationship_1 only). parent_email_2 / parent_relationship_2 are
  always left blank, and no student ever gets more than one email.

This script only ADDS rows to an already-downloaded Weave import
template — it never builds a workbook from scratch when that template
exists, and it never touches any sheet other than the student data
sheet (so the signed `_import_metadata` sheet and the `Instructions`
sheet are left completely untouched).

Requires: pip install faker openpyxl
"""

import os
import random

from faker import Faker
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

# Nigerian locale gives more realistic Nigerian first/last names
fake = Faker("en_NG")

# ---- Config ----------------------------------------------------------
NUM_ROWS = 1000  # how many fake students to generate

# IMPORTANT: on Windows, this must be a raw string (the r prefix below)
# or forward slashes. Without it, backslash sequences like \U, \t, \n
# get interpreted as escape codes and can crash the script or silently
# mangle the path.
OUTPUT_FILE = r"c:\Users\taiwo\Downloads\students_import_template (6).xlsx"

# Arms available — A, B, C for every level.
ARMS_BY_LEVEL = {
    "JSS1": ["A", "B", "C"],
    "JSS2": ["A", "B", "C"],
    "JSS3": ["A", "B", "C"],
    "SS1": ["A", "B", "C"],
    "SS2": ["A", "B", "C"],
    "SS3": ["A", "B", "C"],
}

# Roughly the age (in years) a student in each level would be
AGE_RANGE_BY_LEVEL = {
    "JSS1": (10, 11),
    "JSS2": (11, 12),
    "JSS3": (12, 13),
    "SS1": (13, 14),
    "SS2": (14, 15),
    "SS3": (15, 16),
}

# Nigerian states (Faker has no built-in "state" provider for en_NG)
STATES_OF_ORIGIN = [
    "Lagos",
    "Ogun",
    "Oyo",
    "Kano",
    "Rivers",
    "Enugu",
    "Kaduna",
    "Abia",
    "Delta",
    "Edo",
    "Anambra",
    "Imo",
    "Osun",
    "Ondo",
    "Plateau",
    "Benue",
    "Cross River",
    "Akwa Ibom",
    "Sokoto",
    "Borno",
]

FIELDNAMES = [
    "first_name",
    "last_name",
    "date_of_birth",
    "gender",
    "level",
    "arm",
    "state_of_origin",
    "parent_email_1",
    "parent_relationship_1",
    "parent_email_2",
    "parent_relationship_2",
]

PARENT_RELATIONSHIPS = ["Father", "Mother", "Guardian"]

# Single real inbox used for every generated parent email, via Gmail's
# "+" alias trick. dfragraid+parent7@gmail.com still delivers to
# dfragraid@gmail.com, but reads as a distinct address.
PARENT_EMAIL_LOCAL, PARENT_EMAIL_DOMAIN = "dfragraid", "gmail.com"


def parent_alias_email(n: int) -> str:
    return f"{PARENT_EMAIL_LOCAL}+parent{n}@{PARENT_EMAIL_DOMAIN}"


def random_dob(level: str) -> str:
    """Return a plausible date of birth (YYYY-MM-DD) for the given level."""
    min_age, max_age = AGE_RANGE_BY_LEVEL[level]
    return fake.date_of_birth(minimum_age=min_age, maximum_age=max_age).isoformat()


def generate_row(level=None, arm=None, parent_email=None):
    """Build one fake student row.

    level/arm can be forced (used to guarantee class coverage); otherwise
    they're chosen at random. If parent_email is given, it's written into
    parent_email_1 with a random relationship; parent_email_2 is always
    left blank — each student gets at most one parent email.
    """
    if level is None:
        level = random.choice(list(ARMS_BY_LEVEL))
    if arm is None:
        arm = random.choice(ARMS_BY_LEVEL[level])

    gender = random.choice(["Male", "Female"])
    first_name = fake.first_name_male() if gender == "Male" else fake.first_name_female()

    if parent_email:
        parent_email_1 = parent_email
        parent_relationship_1 = random.choice(PARENT_RELATIONSHIPS)
    else:
        parent_email_1 = ""
        parent_relationship_1 = ""

    return {
        "first_name": first_name,
        "last_name": fake.last_name(),
        "date_of_birth": random_dob(level),
        "gender": gender,
        "level": level,
        "arm": arm,
        "state_of_origin": random.choice(STATES_OF_ORIGIN),
        "parent_email_1": parent_email_1,
        "parent_relationship_1": parent_relationship_1,
        "parent_email_2": "",
        "parent_relationship_2": "",
    }


def build_rows():
    """Generate NUM_ROWS rows, guaranteeing at least one parent email per
    class (level + arm), and no more than one parent email per student."""
    all_classes = [
        (level, arm) for level in ARMS_BY_LEVEL for arm in ARMS_BY_LEVEL[level]
    ]  # 6 levels x 3 arms = 18 classes

    if NUM_ROWS < len(all_classes):
        raise ValueError(
            f"NUM_ROWS ({NUM_ROWS}) is smaller than the number of classes "
            f"({len(all_classes)}); can't guarantee one covered student per "
            f"class. Increase NUM_ROWS."
        )

    rows = []

    # One explicitly covered student per class, each with a unique alias email.
    for i, (level, arm) in enumerate(all_classes, start=1):
        rows.append(generate_row(level=level, arm=arm, parent_email=parent_alias_email(i)))

    # Fill the rest randomly, with no parent email at all.
    for _ in range(NUM_ROWS - len(all_classes)):
        rows.append(generate_row())

    random.shuffle(rows)  # so the covered rows aren't suspiciously grouped at the top
    return rows


def main():
    rows = build_rows()

    if os.path.exists(OUTPUT_FILE):
        # Open the existing template as-is. This preserves its formatting,
        # column order, any other sheets, data validation, etc. We only
        # add new rows underneath whatever is already there.
        wb = load_workbook(OUTPUT_FILE)
        ws = wb.active

        # Map "column name" -> column index by reading the existing header
        # row, so this works regardless of what order the template's
        # columns are actually in.
        header_row = 1
        col_map = {}
        for cell in ws[header_row]:
            if cell.value is not None:
                col_map[str(cell.value).strip().lower()] = cell.column

        missing = [name for name in FIELDNAMES if name.lower() not in col_map]
        if missing:
            raise ValueError(
                f"The downloaded template's header row is missing these "
                f"expected column(s): {missing}. This usually means the "
                f"template predates the current level/arm import "
                f"contract. Download a fresh template from Weave and rerun "
                f"this script against it — do NOT rename or add columns "
                f"in the existing file to work around this."
            )

        start_row = ws.max_row + 1
        for offset, row_data in enumerate(rows):
            row_idx = start_row + offset
            for name, value in row_data.items():
                ws.cell(row=row_idx, column=col_map[name.lower()], value=value)

        wb.save(OUTPUT_FILE)
        print(f"Appended {len(rows)} fake student rows to the existing file: {OUTPUT_FILE}")

    else:
        # File doesn't exist yet — create it fresh with a header row.
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"

        ws.append(FIELDNAMES)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for row_data in rows:
            ws.append([row_data[name] for name in FIELDNAMES])

        widths = [14, 14, 14, 10, 8, 8, 16, 26, 20, 22, 20]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

        wb.save(OUTPUT_FILE)
        print(f"File didn't exist yet — created it with {len(rows)} fake student rows: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()