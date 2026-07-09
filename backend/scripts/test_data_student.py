
"""
Fills students_import_template.xlsx with fake student data.
Requires: pip install openpyxl faker
"""

from pathlib import Path
from random import choice, randint

from faker import Faker
from openpyxl import load_workbook


FILE_PATH = Path(r"C:\Users\taiwo\Downloads\students_import_template.xlsx")

ROW_COUNT = 1500

HEADERS = [
    "first_name",
    "last_name",
    "date_of_birth",
    "gender",
    "class_name",
    "class_arm",
    "state_of_origin",
]

# IMPORTANT:
# These class names/arms must already exist in your backend.
# If they do not exist, dry-run will fail with class_not_found.
CLASS_OPTIONS = [
    ("JSS1", "A"),
    ("JSS1", "B"),
    ("JSS2", "A"),
    ("JSS2", "B"),
    ("JSS3", "A"),
    ("SS1", "A"),
    ("SS2", "A"),
    ("SS3", "A"),
]

NIGERIAN_STATES = [
    "Abia",
    "Adamawa",
    "Akwa Ibom",
    "Anambra",
    "Bauchi",
    "Bayelsa",
    "Benue",
    "Borno",
    "Cross River",
    "Delta",
    "Ebonyi",
    "Edo",
    "Ekiti",
    "Enugu",
    "FCT",
    "Gombe",
    "Imo",
    "Jigawa",
    "Kaduna",
    "Kano",
    "Katsina",
    "Kebbi",
    "Kogi",
    "Kwara",
    "Lagos",
    "Nasarawa",
    "Niger",
    "Ogun",
    "Ondo",
    "Osun",
    "Oyo",
    "Plateau",
    "Rivers",
    "Sokoto",
    "Taraba",
    "Yobe",
    "Zamfara",
]

fake = Faker()


def build_date_of_birth() -> str:
    """Build a realistic student date of birth in YYYY-MM-DD format."""

    birth_year = randint(2007, 2018)
    birth_month = randint(1, 12)
    birth_day = randint(1, 28)

    return f"{birth_year}-{birth_month:02d}-{birth_day:02d}"


def build_row() -> list[str]:
    """Build one student row matching the backend student import template."""

    class_name, class_arm = choice(CLASS_OPTIONS)

    return [
        fake.first_name(),
        fake.last_name(),
        build_date_of_birth(),
        choice(["male", "female"]),
        class_name,
        class_arm,
        choice(NIGERIAN_STATES),
    ]


def clear_existing_data(sheet) -> None:
    """Clear old data rows while keeping the template metadata intact."""

    if sheet.max_row <= 1:
        return

    sheet.delete_rows(2, sheet.max_row - 1)


def ensure_headers(sheet) -> None:
    """Ensure the visible sheet headers match the student import template."""

    existing_headers = [
        sheet.cell(row=1, column=col_index).value
        for col_index in range(1, len(HEADERS) + 1)
    ]

    if existing_headers != HEADERS:
        for col_index, header in enumerate(HEADERS, start=1):
            sheet.cell(row=1, column=col_index, value=header)


def write_rows(sheet) -> None:
    """Write fake student rows into the template."""

    for row_offset in range(ROW_COUNT):
        row_number = row_offset + 2
        row_data = build_row()

        for col_index, value in enumerate(row_data, start=1):
            sheet.cell(row=row_number, column=col_index, value=value)


def main() -> None:
    workbook = load_workbook(FILE_PATH)
    sheet = workbook.active

    ensure_headers(sheet)
    clear_existing_data(sheet)
    write_rows(sheet)

    workbook.save(FILE_PATH)
    print(f"Wrote {ROW_COUNT} student rows to {FILE_PATH}")


if __name__ == "__main__":
    main()