
"""
Fills teachers_import_template.xlsx with fake teacher data.
Requires: pip install openpyxl faker
"""

from pathlib import Path

from faker import Faker
from openpyxl import load_workbook


FILE_PATH = Path(r"C:\Users\taiwo\Downloads\teachers_import_template.xlsx")

ROW_COUNT = 50

BASE_EMAIL_USER = "dfragraid"
BASE_EMAIL_DOMAIN = "gmail.com"

HEADERS = [
    "email",
    "first_name",
    "last_name",
    "staff_id",
    "qualification",
    "specialization",
]

QUALIFICATIONS = [
    "B.Ed",
    "B.Sc",
    "M.Ed",
    "M.Sc",
    "NCE",
    "PGDE",
]

SPECIALIZATIONS = [
    "Mathematics",
    "English Language",
    "Basic Science",
    "Computer Studies",
    "Biology",
    "Chemistry",
    "Physics",
    "Economics",
    "Government",
    "Literature",
    "Civic Education",
    "Business Studies",
]


fake = Faker()


def build_staff_id(index: int) -> str:
    """Build a unique staff ID for each teacher row."""

    return f"TCH-{index:04d}"


def build_email(index: int) -> str:
    """Build a unique test email address."""

    return f"{BASE_EMAIL_USER}+teacher{index}@{BASE_EMAIL_DOMAIN}"


def build_row(index: int) -> list[str]:
    """Build one teacher import row matching the backend template."""

    return [
        build_email(index),
        fake.first_name(),
        fake.last_name(),
        build_staff_id(index),
        fake.random_element(QUALIFICATIONS),
        fake.random_element(SPECIALIZATIONS),
    ]


def clear_existing_data(sheet) -> None:
    """Clear old data rows while keeping the template header and metadata intact."""

    if sheet.max_row <= 1:
        return

    sheet.delete_rows(2, sheet.max_row - 1)


def ensure_headers(sheet) -> None:
    """Ensure the visible sheet headers match the teacher import template."""

    existing_headers = [
        sheet.cell(row=1, column=col_index).value
        for col_index in range(1, len(HEADERS) + 1)
    ]

    if existing_headers != HEADERS:
        for col_index, header in enumerate(HEADERS, start=1):
            sheet.cell(row=1, column=col_index, value=header)


def write_rows(sheet) -> None:
    """Write fake teacher rows into the template."""

    for row_offset in range(ROW_COUNT):
        row_number = row_offset + 2
        row_data = build_row(row_offset + 1)

        for col_index, value in enumerate(row_data, start=1):
            sheet.cell(row=row_number, column=col_index, value=value)


def main() -> None:
    workbook = load_workbook(FILE_PATH)
    sheet = workbook.active

    ensure_headers(sheet)
    clear_existing_data(sheet)
    write_rows(sheet)

    workbook.save(FILE_PATH)
    print(f"Wrote {ROW_COUNT} teacher rows to {FILE_PATH}")


if __name__ == "__main__":
    main()